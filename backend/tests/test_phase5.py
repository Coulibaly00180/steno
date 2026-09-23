"""Phase 5: speakers (diarization, renaming, merging) and DOCX/PDF reports, translated subtitles."""
import io
from datetime import datetime, timezone

import numpy as np
import pytest
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker

from app import diarization, main, worker
from app.db import Base
from app.diarization import Cluster, Turn, assign_speakers, cluster_speakers, number_by_appearance
from app.exports import translated_cues
from app.models import Passage, ProcessingJob, Speaker, Summary, TranscriptSegment, Video, VideoIndex
from app.reports import Block, Report, build_report, inline_runs, markdown_blocks, pdf_story, render_docx, render_pdf

NOW = datetime(2026, 9, 23, 12, 0, tzinfo=timezone.utc)
MEETING = [
    (0.0, 6.0, "Bonjour à tous, on commence la réunion budget."),
    (6.0, 12.0, "Je propose de valider le budget de 40 000 euros."),
    (12.0, 18.0, "D'accord, budget validé. Je prépare le devis pour vendredi."),
    (18.0, 24.0, "Parfait, merci. Fin de la réunion."),
]


@pytest.fixture
def db_session(monkeypatch, tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'p5.db'}")

    @event.listens_for(engine, "connect")
    def enable_foreign_keys(dbapi_connection, _):
        dbapi_connection.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine, expire_on_commit=False)
    monkeypatch.setattr(main, "SessionLocal", session)
    monkeypatch.setattr(worker, "SessionLocal", session)
    monkeypatch.setattr(main.settings, "data_dir", tmp_path)
    (tmp_path / "audio").mkdir()
    return session


def add_meeting(session, video_id="v", *, status="COMPLETED", translated=None, summary=None, diarize=False):
    with session() as db:
        db.add(Video(
            id=video_id, filename="m.mp4", original_filename="Réunion budget.mp4", path="/nowhere/m.mp4",
            duration_seconds=24, size_bytes=1, status=status, detected_language="fr", created_at=NOW,
            transcript_text="\n".join(f"[00:00:{int(s):02d}] {t}" for s, _, t in MEETING), translated_text=translated,
            diarize=diarize,
        ))
        db.flush()
        db.add_all([TranscriptSegment(video_id=video_id, start_seconds=s, end_seconds=e, text=t) for s, e, t in MEETING])
        if summary:
            db.add(Summary(id=f"s-{video_id}", video_id=video_id, content_markdown=summary, model="m", created_at=NOW))
        db.commit()


def fake_turns(monkeypatch, turns=None):
    """Alice speaks at 0-6 and 18-24, Bob at 6-18."""
    turns = turns or [Turn(0, 6, 1), Turn(6, 18, 2), Turn(18, 24, 1)]
    calls = []

    def fake(path, *, num_speakers=None, on_progress=None):
        calls.append((path.name, num_speakers))
        if on_progress:
            on_progress(24)
        return turns

    monkeypatch.setattr(worker, "diarize", fake)
    return calls


# --- diarization pieces ------------------------------------------------------------------

def voice(seed: int, noise: float = 0.05, size: int = 16) -> np.ndarray:
    base = np.random.default_rng(seed).normal(size=size)
    vector = base + np.random.default_rng(seed * 100 + int(noise * 1000)).normal(scale=noise, size=size)
    return (vector / np.linalg.norm(vector)).astype(np.float32)


def test_clusters_of_one_voice_merge_and_fragments_join_a_speaker():
    clusters = [
        Cluster([(0, 60)], voice(1, 0.05), 60), Cluster([(60, 120)], voice(1, 0.1), 60),
        Cluster([(120, 180)], voice(2, 0.05), 60), Cluster([(180, 240)], voice(2, 0.1), 60),
        Cluster([(240, 300)], voice(3, 0.05), 60),
        # 3 seconds of an unclear voice print: a fragment, not a fourth speaker.
        Cluster([(300, 303)], voice(9, 0.05), 3),
    ]
    labels = cluster_speakers(clusters)
    assert labels[0] == labels[1] and labels[2] == labels[3]
    assert len({labels[0], labels[2], labels[4]}) == 3
    assert len(set(labels)) == 3


def test_a_given_number_of_speakers_is_honoured():
    clusters = [Cluster([(i * 60, i * 60 + 60)], voice(i), 60) for i in range(5)]
    assert len(set(cluster_speakers(clusters, 3))) == 3
    assert len(set(cluster_speakers(clusters, 8))) == 5  # never more groups than clusters


def test_speakers_are_numbered_by_first_appearance_and_turns_joined():
    turns = number_by_appearance([(10, 12, 7), (0, 5, 3), (5.2, 9, 3), (12.1, 14, 7)])
    assert turns == [Turn(0, 9, 1), Turn(10, 14, 2)]


def test_each_segment_goes_to_the_speaker_talking_most():
    turns = [Turn(0, 6, 1), Turn(6, 18, 2), Turn(18, 24, 1)]
    assert assign_speakers([(0, 6), (5, 12), (12, 18), (18.5, 20), (40, 41)], turns) == [1, 2, 2, 1, 1]
    assert assign_speakers([(0, 1)], []) == [None]


# --- worker ------------------------------------------------------------------------------

def test_diarized_transcript_names_every_line(db_session, monkeypatch, tmp_path):
    add_meeting(db_session)
    (tmp_path / "audio" / "v.wav").write_bytes(b"wav")
    calls = fake_turns(monkeypatch)
    with db_session() as db:
        db.get(Video, "v").num_speakers = 2
        db.commit()

    assert worker.diarize_video("v") == 2
    assert calls == [("v.wav", 2)]
    with db_session() as db:
        video = db.get(Video, "v")
        assert [s.position for s in video.speakers] == [1, 2]
        assert video.transcript_text.splitlines()[1] == "[00:00:06] Intervenant 2 : Je propose de valider le budget de 40 000 euros."
        assert video.transcript_text.splitlines()[0].startswith("[00:00:00] Intervenant 1 : Bonjour")


def test_a_second_diarization_replaces_the_speakers(db_session, monkeypatch, tmp_path):
    add_meeting(db_session)
    (tmp_path / "audio" / "v.wav").write_bytes(b"wav")
    fake_turns(monkeypatch)
    worker.diarize_video("v")
    fake_turns(monkeypatch, [Turn(0, 24, 1)])
    assert worker.diarize_video("v") == 1
    with db_session() as db:
        assert db.query(Speaker).count() == 1


def test_pipeline_diarization_failure_keeps_the_video(db_session, monkeypatch, tmp_path):
    add_meeting(db_session, status="PROCESSING", diarize=True)
    (tmp_path / "audio" / "v.wav").write_bytes(b"wav")
    with db_session() as db:
        db.add(ProcessingJob(id="job", video_id="v", status="RUNNING", stage="TRANSCRIBED", progress=50))
        db.commit()
    monkeypatch.setattr(worker, "diarize", lambda *_, **__: (_ for _ in ()).throw(RuntimeError("model missing")))
    worker._diarize_in_pipeline("job", "v", 24)
    with db_session() as db:
        assert db.get(Video, "v").diarization_error == "model missing"
        assert db.get(ProcessingJob, "job").status == "RUNNING"


def test_diarize_job_flags_the_summary_and_reindexes(db_session, monkeypatch, tmp_path):
    add_meeting(db_session, summary="# Décisions\n- Budget validé")
    (tmp_path / "audio" / "v.wav").write_bytes(b"wav")
    fake_turns(monkeypatch)
    queued = []
    monkeypatch.setattr(worker, "enqueue_index_job", queued.append)
    with db_session() as db:
        db.add(ProcessingJob(id="d", video_id="v", kind="DIARIZE", status="QUEUED", stage="QUEUED", progress=0))
        db.add(VideoIndex(video_id="v", status="READY", model="m", transcript_hash="h", passages=1, updated_at=NOW))
        db.commit()

    worker.run_diarize("d")

    with db_session() as db:
        video = db.get(Video, "v")
        assert db.get(ProcessingJob, "d").status == "COMPLETED"
        assert video.status == "COMPLETED" and video.transcript_edited_at is not None
        assert db.get(VideoIndex, "v").status == "STALE"
    assert queued == ["v"]
    assert "Intervenant 2 : Je propose" in (tmp_path / "exports" / "v" / "transcript.srt").read_text(encoding="utf-8")


def test_diarize_job_failure_only_fails_the_job(db_session, monkeypatch):
    add_meeting(db_session)
    with db_session() as db:
        db.add(ProcessingJob(id="d", video_id="v", kind="DIARIZE", status="QUEUED", stage="QUEUED", progress=0))
        db.commit()
    worker.run_diarize("d")  # no audio, no source file
    with db_session() as db:
        assert db.get(ProcessingJob, "d").status == "FAILED"
        assert db.get(ProcessingJob, "d").error == worker.ERROR_SOURCE_NOT_FOUND
        assert db.get(Video, "v").status == "COMPLETED"


def test_passages_carry_the_speaker_names(db_session, monkeypatch, tmp_path):
    add_meeting(db_session)
    (tmp_path / "audio" / "v.wav").write_bytes(b"wav")
    fake_turns(monkeypatch)
    worker.diarize_video("v")
    worker.index_video("v")
    with db_session() as db:
        assert "Intervenant 2 : Je propose" in db.query(Passage).first().text


# --- API -----------------------------------------------------------------------------------

@pytest.fixture
def diarized(db_session, monkeypatch, tmp_path):
    add_meeting(db_session, summary="# Décisions\n- Budget validé [00:00:12]")
    (tmp_path / "audio" / "v.wav").write_bytes(b"wav")
    fake_turns(monkeypatch)
    worker.diarize_video("v")
    monkeypatch.setattr(main, "enqueue_index_job", lambda video_id: None)
    with db_session() as db:
        return [s.id for s in db.get(Video, "v").speakers]


def test_detail_lists_speakers_with_their_speaking_time(client, diarized):
    body = client.get("/videos/v").json()
    assert [(s["label"], s["seconds"], s["share"]) for s in body["speakers"]] == [("Intervenant 1", 12.0, 0.5), ("Intervenant 2", 12.0, 0.5)]
    assert [s["speaker_id"] for s in body["segments"]] == [diarized[0], diarized[1], diarized[1], diarized[0]]


def test_renaming_a_speaker_rewrites_the_transcript(client, db_session, diarized):
    response = client.put(f"/videos/v/speakers/{diarized[1]}", json={"name": "  Marie   Curie "})
    assert response.status_code == 200
    assert response.json()["speakers"][1]["label"] == "Marie Curie"
    body = client.get("/videos/v").json()
    assert "[00:00:06] Marie Curie : Je propose" in body["transcript_text"]
    assert body["summary_outdated"] is True
    # Another speaker cannot take the same name: merging is the way.
    assert client.put(f"/videos/v/speakers/{diarized[0]}", json={"name": "marie curie"}).status_code == 409
    # An empty name brings back « Intervenant 2 ».
    client.put(f"/videos/v/speakers/{diarized[1]}", json={"name": ""})
    assert "Intervenant 2 : Je propose" in client.get("/videos/v").json()["transcript_text"]


def test_merging_speakers_moves_their_lines(client, db_session, diarized):
    response = client.post(f"/videos/v/speakers/{diarized[1]}/merge", json={"into": diarized[0]})
    assert response.status_code == 200
    assert [s["label"] for s in response.json()["speakers"]] == ["Intervenant 1"]
    assert "Intervenant 2" not in client.get("/videos/v").json()["transcript_text"]
    assert client.post(f"/videos/v/speakers/{diarized[0]}/merge", json={"into": diarized[0]}).status_code == 422
    assert client.post("/videos/v/speakers/999/merge", json={"into": diarized[0]}).status_code == 404


def test_a_line_can_change_speaker_or_lose_it(client, db_session, diarized):
    with db_session() as db:
        segment_id = db.query(TranscriptSegment).filter_by(start_seconds=6.0).one().id
    assert client.patch(f"/videos/v/segments/{segment_id}", json={"speaker_id": diarized[0]}).json()["speaker_id"] == diarized[0]
    assert "[00:00:06] Intervenant 1 : Je propose" in client.get("/videos/v").json()["transcript_text"]
    assert client.patch(f"/videos/v/segments/{segment_id}", json={"speaker_id": None}).json()["speaker_id"] is None
    assert "[00:00:06] Je propose" in client.get("/videos/v").json()["transcript_text"]
    assert client.patch(f"/videos/v/segments/{segment_id}", json={"speaker_id": 999}).status_code == 404
    # Text only: the speaker stays.
    client.patch(f"/videos/v/segments/{segment_id}", json={"speaker_id": diarized[1]})
    assert client.patch(f"/videos/v/segments/{segment_id}", json={"text": "Je propose 45 000 euros."}).json()["speaker_id"] == diarized[1]


def test_detecting_speakers_queues_a_job(client, db_session, monkeypatch):
    add_meeting(db_session)
    from tests.test_upload import SuccessfulQueue

    # The media were deleted to save space (n°14): nothing to listen to.
    assert client.post("/videos/v/speakers/detect", json={}).status_code == 409
    main.settings.audio_dir.mkdir(parents=True, exist_ok=True)
    (main.settings.audio_dir / "v.wav").write_bytes(b"RIFF")

    monkeypatch.setattr(main, "Queue", SuccessfulQueue)
    response = client.post("/videos/v/speakers/detect", json={"num_speakers": 4})
    assert response.status_code == 200 and response.json()["kind"] == "DIARIZE"
    with db_session() as db:
        video = db.get(Video, "v")
        assert video.diarize and video.num_speakers == 4
    # The job is the video's current job: a second request waits for it.
    assert client.post("/videos/v/speakers/detect", json={}).status_code == 409
    assert client.post("/videos/v/speakers/detect", json={"num_speakers": 0}).status_code == 422


def test_upload_validates_the_number_of_speakers(client, upload_environment, monkeypatch):
    from tests.test_upload import SuccessfulQueue

    session, _ = upload_environment
    monkeypatch.setattr(main, "Queue", SuccessfulQueue)
    files = {"file": ("clip.mp4", b"media", "video/mp4")}
    assert client.post("/videos", files=files, data={"diarize": "true", "num_speakers": "21"}).status_code == 422
    response = client.post("/videos", files=files, data={"diarize": "true", "num_speakers": "4"})
    assert response.status_code == 200
    with session() as db:
        video = db.get(Video, response.json()["video_id"])
        assert video.diarize is True and video.num_speakers == 4


# --- exports -------------------------------------------------------------------------------

def test_translated_subtitles_follow_the_source_timing():
    translated = "[00:00:00] Hello everyone.\n[00:00:06] I propose to approve\nthe budget.\n[00:00:40] End."
    cues = translated_cues(translated, [(0.0, 5.5), (6.0, 11.0), (40.0, 42.0)])
    assert cues == [(0.0, 5.5, "Hello everyone."), (6.0, 11.0, "I propose to approve the budget."), (40.0, 42.0, "End.")]
    # Without a matching segment, a cue lasts until the next one, 7 s at most.
    assert translated_cues("[00:00:01] A\n[00:00:03] B", []) == [(1.0, 3.0, "A"), (3.0, 10.0, "B")]


def test_translation_subtitle_files_are_written(db_session, tmp_path):
    add_meeting(db_session, translated="[00:00:00] Hello everyone.\n[00:00:06] I propose a budget.")
    with db_session() as db:
        worker.write_exports(db, "v")
    srt = (tmp_path / "exports" / "v" / "translation.srt").read_text(encoding="utf-8")
    assert srt.startswith("1\n00:00:00,000 --> 00:00:06,000\nHello everyone.")
    assert (tmp_path / "exports" / "v" / "translation.vtt").read_text(encoding="utf-8").startswith("WEBVTT")


def test_markdown_blocks_and_bold_runs():
    blocks = markdown_blocks("# Décisions\n- **Budget** validé\n1. Devis\n\nUn paragraphe\nsur deux lignes.\n---")
    assert [(b.kind, b.text, b.level) for b in blocks] == [
        ("heading", "Décisions", 1), ("bullet", "**Budget** validé", 0), ("numbered", "Devis", 0),
        ("paragraph", "Un paragraphe sur deux lignes.", 0),
    ]
    assert inline_runs("**Marie** valide le **budget**.") == [("Marie", True), (" valide le ", False), ("budget", True), (".", False)]


def test_docx_report_has_speakers_decisions_and_transcript(client, db_session, diarized):
    client.put(f"/videos/v/speakers/{diarized[1]}", json={"name": "Marie"})
    with db_session() as db:
        db.get(Summary, "s-v").content_markdown = "# Décisions\n- **Budget** validé par Marie\n# Actions\n- Marie : devis pour vendredi"
        db.commit()
    response = client.get("/videos/v/exports/report.docx")
    assert response.status_code == 200
    assert "compte-rendu.docx" in response.headers["content-disposition"]
    from docx import Document

    document = Document(io.BytesIO(response.content))
    text = "\n".join(p.text for p in document.paragraphs)
    tables = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "Budget validé par Marie" in text and "Marie : devis pour vendredi" in text
    assert "Marie : Je propose de valider le budget de 40 000 euros." in text
    assert "Marie" in tables and "50 %" in tables
    bold = [run.text for p in document.paragraphs for run in p.runs if run.bold]
    assert "Budget" in bold


def test_pdf_report_is_generated(client, db_session, diarized):
    response = client.get("/videos/v/exports/report.pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF") and len(response.content) > 2000
    report = build_report(db_session(), db_session().get(Video, "v"))
    assert render_pdf(report).startswith(b"%PDF")
    assert render_docx(report)[:2] == b"PK"


def test_pdf_lists_keep_their_items():
    """Every bullet list of the PDF used to come out empty (a list cleared after use)."""
    from reportlab.platypus import ListFlowable

    report = Report(title="R", details=[("Fichier", "r.mp4")], summary=[
        Block("heading", "Décisions", 1), Block("bullet", "Budget validé"), Block("bullet", "Devis vendredi"),
        Block("heading", "Actions", 1), Block("bullet", "Marie : devis"),
    ])
    lists = [flowable for flowable in pdf_story(report) if isinstance(flowable, ListFlowable)]
    assert [len(flowable._flowables) for flowable in lists] == [2, 1]


def test_reports_need_a_processed_video(client, db_session):
    add_meeting(db_session, status="PROCESSING")
    assert client.get("/videos/v/exports/report.docx").status_code == 404
    assert client.get("/videos/missing/exports/report.pdf").status_code == 404


def test_models_directory_check(tmp_path, monkeypatch):
    monkeypatch.setattr(diarization.settings, "diarization_models_dir", tmp_path)
    assert diarization.models_available() is False
    (tmp_path / "segmentation.onnx").write_bytes(b"x")
    (tmp_path / "embedding.onnx").write_bytes(b"x")
    assert diarization.models_available() is True
