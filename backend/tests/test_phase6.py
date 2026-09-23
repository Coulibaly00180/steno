"""Finishing touches: saved conversations and interrupted answers, translated reports, snippets."""
import asyncio

import pytest

from app import main, worker
from app.models import LibraryConversation, LibraryMessage, VideoChatMessage
from app.schemas import ChatQuestion, LibraryQuestion
from tests.test_phase4 import add_video, db_session, sse_events  # noqa: F401  (fixture reused)


def two_indexed_videos(session):
    add_video(session, "budget", [(float(i * 10), float(i * 10 + 10), f"Le budget annuel, point {i}.") for i in range(6)], name="Budget.mp4")
    add_video(session, "velo", [(0.0, 10.0, "Le conseil achète des vélos électriques.")], name="Conseil.mp4")
    worker.index_video("budget")
    worker.index_video("velo")


def fake_stream(monkeypatch, *pieces, then=None):
    async def stream(prompt, **kwargs):
        for piece in pieces:
            yield piece
        if then:
            await then()

    monkeypatch.setattr(main, "stream_chat", stream)


# --- n°19: conversations are kept -------------------------------------------------------------

def test_a_library_conversation_is_saved_listed_reopened_and_deleted(client, db_session, monkeypatch):
    two_indexed_videos(db_session)
    fake_stream(monkeypatch, "Le budget est voté ", "[1].")
    first = sse_events(client.post("/library/chat/stream", json={
        "question": "Que dit-on du budget ?", "video_ids": ["budget", "velo"], "scope": "tag « Finance »",
    }).text)
    conversation_id = first[0][1]["conversation_id"]
    client.post("/library/chat/stream", json={"question": "Et les vélos ?", "conversation_id": conversation_id})

    listing = client.get("/library/conversations").json()
    assert [(c["id"], c["title"], c["scope"], c["video_count"]) for c in listing] == [
        (conversation_id, "Que dit-on du budget ?", "tag « Finance »", 2)
    ]
    detail = client.get(f"/library/conversations/{conversation_id}").json()
    assert [(m["role"], m["content"]) for m in detail["messages"]] == [
        ("user", "Que dit-on du budget ?"), ("assistant", "Le budget est voté [1]."),
        ("user", "Et les vélos ?"), ("assistant", "Le budget est voté [1]."),
    ]
    assert detail["video_ids"] == ["budget", "velo"]
    # The cited passages are kept with the answer: the [1] still opens its video.
    assert detail["messages"][1]["sources"][0]["video_id"] in {"budget", "velo"}

    assert client.delete(f"/library/conversations/{conversation_id}").json() == {"deleted": True}
    assert client.get(f"/library/conversations/{conversation_id}").status_code == 404
    with db_session() as db:
        assert db.query(LibraryMessage).count() == 0


def test_a_follow_up_needs_an_existing_conversation(client, db_session):
    two_indexed_videos(db_session)
    response = client.post("/library/chat/stream", json={"question": "Q ?", "conversation_id": "missing"})
    assert response.status_code == 404
    assert client.post("/library/chat/stream", json={"question": "Q ?"}).status_code == 422


def test_a_failed_library_answer_keeps_its_partial_text(client, db_session, monkeypatch):
    two_indexed_videos(db_session)

    async def failing(prompt, **kwargs):
        yield "Le budget"
        raise ConnectionError("stopped")

    monkeypatch.setattr(main, "stream_chat", failing)
    events = sse_events(client.post("/library/chat/stream", json={"question": "Budget ?", "video_ids": ["budget"]}).text)
    assert events[-1][1]["saved"] is True
    with db_session() as db:
        answer = db.query(LibraryMessage).filter_by(role="assistant").one()
        assert (answer.content, answer.interrupted) == ("Le budget", True)


def test_a_library_conversation_without_any_answer_is_not_kept(client, db_session, monkeypatch):
    two_indexed_videos(db_session)

    async def failing(prompt, **kwargs):
        raise ConnectionError("down")
        yield  # pragma: no cover

    monkeypatch.setattr(main, "stream_chat", failing)
    events = sse_events(client.post("/library/chat/stream", json={"question": "Budget ?", "video_ids": ["budget"]}).text)
    assert events[-1][1]["saved"] is False
    with db_session() as db:
        assert db.query(LibraryConversation).count() == 0


# --- leaving the page keeps what was written ----------------------------------------------------

def test_leaving_a_video_chat_mid_answer_keeps_the_written_part(db_session, monkeypatch):
    add_video(db_session, "v", [(0.0, 5.0, "Le budget est validé.")])

    async def forever():
        await asyncio.sleep(30)

    fake_stream(monkeypatch, "Le budget ", "est validé", then=forever)

    async def read_two_pieces_then_leave():
        response = await main.stream_video_answer("v", ChatQuestion(question="Et le budget ?"))
        stream = response.body_iterator
        await stream.__anext__()
        await stream.__anext__()
        await stream.aclose()  # the browser closed the connection

    asyncio.run(read_two_pieces_then_leave())
    with db_session() as db:
        messages = db.query(VideoChatMessage).order_by(VideoChatMessage.created_at).all()
        assert [(m.role, m.content, m.interrupted) for m in messages] == [
            ("user", "Et le budget ?", False), ("assistant", "Le budget est validé", True),
        ]


def test_leaving_a_library_answer_mid_way_keeps_it_and_the_conversation(db_session, monkeypatch):
    two_indexed_videos(db_session)

    async def forever():
        await asyncio.sleep(30)

    fake_stream(monkeypatch, "Selon [1]", then=forever)

    async def read_then_leave():
        response = await main.stream_library_answer(LibraryQuestion(question="Budget ?", video_ids=["budget"]))
        stream = response.body_iterator
        await stream.__anext__()  # sources
        await stream.__anext__()  # first piece
        await stream.aclose()

    asyncio.run(read_then_leave())
    with db_session() as db:
        assert db.query(LibraryConversation).count() == 1
        answer = db.query(LibraryMessage).filter_by(role="assistant").one()
        assert (answer.content, answer.interrupted) == ("Selon [1]", True)


def test_leaving_before_any_word_keeps_nothing(db_session, monkeypatch):
    two_indexed_videos(db_session)

    async def forever():
        await asyncio.sleep(30)

    fake_stream(monkeypatch, then=forever)

    async def read_sources_then_leave():
        response = await main.stream_library_answer(LibraryQuestion(question="Budget ?", video_ids=["budget"]))
        stream = response.body_iterator
        await stream.__anext__()
        await stream.aclose()

    asyncio.run(read_sources_then_leave())
    with db_session() as db:
        assert db.query(LibraryConversation).count() == 0


@pytest.mark.parametrize("payload", [{"question": "Q ?", "video_ids": ["x"] * 501}])
def test_library_question_limits(client, payload):
    assert client.post("/library/chat/stream", json=payload).status_code == 422


# --- n°20: translated reports ---------------------------------------------------------------------

def translated_meeting(session, monkeypatch, tmp_path):
    """The 4-line meeting of test_phase5, diarized, translated, with a summary."""
    from tests.test_phase5 import MEETING, Turn, fake_turns  # noqa: F401
    from app.models import Summary, TranscriptSegment, Video
    from datetime import datetime, timezone

    with session() as db:
        db.add(Video(
            id="m", filename="m.mp4", original_filename="Réunion budget.mp4", path="/x", duration_seconds=24, size_bytes=1,
            status="COMPLETED", detected_language="fr", target_language="anglais", created_at=datetime(2026, 9, 23, tzinfo=timezone.utc),
            transcript_text="x",
            translated_text=(
                "[00:00:00] Intervenant 1 : Hello everyone, let's start the budget meeting.\n"
                "[00:00:06] Intervenant 2 : I propose to approve the budget: 40,000 euros.\n"
                "wrapped continuation line\n"
                "[00:00:12] Somebody said : this is not a speaker.\n"
            ),
        ))
        db.flush()
        db.add_all([TranscriptSegment(video_id="m", start_seconds=s, end_seconds=e, text=t) for s, e, t in MEETING])
        db.add(Summary(id="s", video_id="m", content_markdown="# Decisions\n- Budget approved", model="m",
                       created_at=datetime(2026, 9, 23, tzinfo=timezone.utc)))
        db.commit()
    (tmp_path / "audio").mkdir(exist_ok=True)
    (tmp_path / "audio" / "m.wav").write_bytes(b"wav")
    fake_turns(monkeypatch)
    worker.diarize_video("m")


def test_translated_rows_keep_only_known_speakers():
    from app.reports import translated_rows

    rows = translated_rows(
        "[00:00:00] Marie : Hello.\n[01:02] Bob : Hi\ncontinued\n[00:00:09] Marie: no space before the colon\n",
        ["Marie"],
    )
    assert rows == [
        ("00:00:00", "Marie", "Hello."),
        ("00:01:02", None, "Bob : Hi continued"),  # Bob is not a speaker of this video
        ("00:00:09", None, "Marie: no space before the colon"),
    ]


def test_report_annex_can_be_original_translation_or_none(client, db_session, monkeypatch, tmp_path):
    import io

    from docx import Document

    translated_meeting(db_session, monkeypatch, tmp_path)

    def paragraphs(query):
        response = client.get(f"/videos/m/exports/report.docx{query}")
        assert response.status_code == 200, response.text
        return response, [p.text for p in Document(io.BytesIO(response.content)).paragraphs], Document(io.BytesIO(response.content))

    _, original, _ = paragraphs("")
    assert any("Bonjour à tous, on commence la réunion budget." in line for line in original)

    response, translated, document = paragraphs("?transcript=translation")
    assert any(line == "[00:00:06] Intervenant 2 : I propose to approve the budget: 40,000 euros. wrapped continuation line" for line in translated)
    assert not any("Bonjour à tous" in line for line in translated)
    assert "compte-rendu%20%28traduit%29.docx" in response.headers["content-disposition"]
    assert any("traduite (anglais)" in cell.text for table in document.tables for row in table.rows for cell in row.cells)

    _, none, _ = paragraphs("?transcript=none")
    assert "Transcription" not in none and not any("Bonjour" in line for line in none)
    assert client.get("/videos/m/exports/report.pdf?transcript=translation").content.startswith(b"%PDF")
    assert client.get("/videos/m/exports/report.pdf?transcript=weird").status_code == 422


def test_report_without_translation_refuses_the_translated_annex(client, db_session, monkeypatch, tmp_path):
    translated_meeting(db_session, monkeypatch, tmp_path)
    with db_session() as db:
        db.get(main.Video, "m").translated_text = None
        db.commit()
    assert client.get("/videos/m/exports/report.docx?transcript=translation").status_code == 404


def test_renaming_or_merging_a_speaker_updates_the_translation_too(client, db_session, monkeypatch, tmp_path):
    translated_meeting(db_session, monkeypatch, tmp_path)
    with db_session() as db:
        ids = [s.id for s in db.get(main.Video, "m").speakers]
    client.put(f"/videos/m/speakers/{ids[1]}", json={"name": "Marie"})
    with db_session() as db:
        text = db.get(main.Video, "m").translated_text
        assert "[00:00:06] Marie : I propose" in text and "Intervenant 2 :" not in text
        assert "[00:00:12] Somebody said : this is not a speaker." in text
    client.post(f"/videos/m/speakers/{ids[1]}/merge", json={"into": ids[0]})
    with db_session() as db:
        text = db.get(main.Video, "m").translated_text
        assert "[00:00:06] Intervenant 1 : I propose" in text and "Marie :" not in text
    srt = (tmp_path / "exports" / "m" / "translation.srt").read_text(encoding="utf-8")
    assert "Intervenant 1 : I propose" in srt


# --- n°17: snippets around the matched words ------------------------------------------------------

def test_snippet_is_clean_located_and_timed():
    from app.main import build_snippet

    fragment = "[00:00:00] Bonjour à tous\n[00:01:30] Le budget annuel est validé par Marie"
    match = fragment.index("budget") + 1
    snippet = build_snippet(fragment, 1, match - 1, ["budget", "valide"], "transcript")
    assert snippet["text"] == "Bonjour à tous … Le budget annuel est validé par Marie"
    highlighted = [snippet["text"][a:b] for a, b in snippet["ranges"]]
    assert highlighted == ["budget", "validé"]  # accent-insensitive, the shown text keeps its accents
    assert snippet["start_seconds"] == 90.0 and snippet["source"] == "transcript"


def test_snippet_cut_inside_a_word_gets_ellipses_and_shifted_ranges():
    from app.main import SNIPPET_LENGTH, build_snippet

    fragment = "rtaine phrase [00:05:00] avec le mot budget au milieu " + "x" * SNIPPET_LENGTH
    snippet = build_snippet(fragment[:SNIPPET_LENGTH], 500, fragment.index("budget"), ["budget"], "translation")
    assert snippet["text"].startswith("… ") and snippet["text"].endswith(" …")
    a, b = snippet["ranges"][0]
    assert snippet["text"][a:b] == "budget"
    assert snippet["start_seconds"] == 300.0


def test_library_search_returns_a_snippet_only_for_a_content_match(client, db_session):
    from app.models import Video

    add_video(db_session, "content", [(0.0, 5.0, "Bonjour à tous"), (90.0, 95.0, "Le budget annuel est validé")], name="Réunion.mp4")
    add_video(db_session, "title", [(0.0, 5.0, "Rien de spécial")], name="Budget 2027.mp4")
    add_video(db_session, "translated", [(0.0, 5.0, "Bonjour")], name="Autre.mp4")
    with db_session() as db:
        db.get(Video, "translated").translated_text = "[00:03:00] The annual budget is approved"
        db.commit()

    rows = {row["id"]: row for row in client.get("/videos", params={"q": "budget"}).json()}
    assert rows["content"]["snippet"]["start_seconds"] == 90.0
    text, ranges = rows["content"]["snippet"]["text"], rows["content"]["snippet"]["ranges"]
    assert [text[a:b] for a, b in ranges] == ["budget"]
    assert rows["title"]["snippet"] is None  # matched on its file name only
    assert rows["translated"]["snippet"]["source"] == "translation"
    assert rows["translated"]["snippet"]["start_seconds"] == 180.0
    assert all(row["snippet"] is None for row in client.get("/videos").json())
