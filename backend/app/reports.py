"""Meeting reports as DOCX and PDF (n°20), built on demand from the database.

Generated at download time rather than stored: a corrected transcript, a
renamed speaker or an edited summary is always in the next download.
"""
import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from .models import SummaryTemplate, Video
from .speakers import speaker_labels, speaking_seconds
from .utils import timestamp

FONT_DIR = Path("/usr/share/fonts/truetype/dejavu")
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_BOLD = re.compile(r"\*\*(.+?)\*\*")


@dataclass(frozen=True)
class Block:
    kind: str  # heading, bullet, numbered, paragraph
    text: str
    level: int = 0


def markdown_blocks(markdown: str) -> list[Block]:
    """The Markdown the summaries use: headings, lists, paragraphs (joined lines)."""
    blocks: list[Block] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(Block("paragraph", " ".join(paragraph)))
            paragraph.clear()

    for raw in markdown.splitlines():
        line = raw.strip()
        if not line or set(line) <= {"-", "*", "_"}:
            flush()
            continue
        if heading := _HEADING.match(line):
            flush()
            blocks.append(Block("heading", heading.group(2).strip("# ").strip(), len(heading.group(1))))
        elif bullet := _BULLET.match(line):
            flush()
            blocks.append(Block("bullet", bullet.group(1)))
        elif numbered := _NUMBERED.match(line):
            flush()
            blocks.append(Block("numbered", numbered.group(1)))
        else:
            paragraph.append(line)
    flush()
    return blocks


def inline_runs(text: str) -> list[tuple[str, bool]]:
    """(text, bold) runs of a line with **bold** parts."""
    runs: list[tuple[str, bool]] = []
    cursor = 0
    for match in _BOLD.finditer(text):
        if match.start() > cursor:
            runs.append((text[cursor:match.start()], False))
        runs.append((match.group(1), True))
        cursor = match.end()
    if cursor < len(text):
        runs.append((text[cursor:], False))
    return runs


@dataclass
class Report:
    title: str
    details: list[tuple[str, str]]
    speakers: list[tuple[str, str, str]] = field(default_factory=list)  # name, time, share
    summary: list[Block] = field(default_factory=list)
    chapters: list[tuple[str, str]] = field(default_factory=list)
    transcript: list[tuple[str, str | None, str]] = field(default_factory=list)  # time, speaker, text


def _duration(seconds: float) -> str:
    minutes = int(round(seconds / 60))
    return f"{minutes // 60} h {minutes % 60:02d}" if minutes >= 60 else f"{minutes} min"


_TRANSLATED_LINE = re.compile(r"^\[(\d{1,2}:\d{2}(?::\d{2})?)\]\s*(.*)$")
TRANSCRIPT_MODES = ("original", "translation", "none")


def translated_rows(translated: str, labels: list[str]) -> list[tuple[str, str | None, str]]:
    """(time, speaker, text) of each translated line; a "Name : " prefix counts only if it is a known speaker."""
    known = {label.casefold(): label for label in labels}
    rows: list[tuple[str, str | None, str]] = []
    for raw in translated.splitlines():
        line = raw.strip()
        if not line:
            continue
        match = _TRANSLATED_LINE.match(line)
        if not match:
            if rows:  # a translated line the model wrapped
                rows[-1] = (rows[-1][0], rows[-1][1], f"{rows[-1][2]} {line}")
            continue
        clock, rest = match.groups()
        head, separator, tail = rest.partition(" : ")
        speaker = known.get(head.strip().casefold()) if separator else None
        rows.append((clock if clock.count(":") == 2 else f"00:{clock}", speaker, (tail if speaker else rest).strip()))
    return rows


def build_report(db, video: Video, *, transcript: str = "original") -> Report:
    summary = video.summaries[-1] if video.summaries else None
    template = db.get(SummaryTemplate, summary.template_id) if summary and summary.template_id else None
    details = [
        ("Fichier", video.original_filename),
        ("Durée", _duration(video.duration_seconds)),
        ("Importé le", video.created_at.astimezone(timezone.utc).strftime("%d/%m/%Y")),
    ]
    if video.detected_language:
        details.append(("Langue parlée", video.detected_language.upper()))
    if transcript == "translation" and video.target_language:
        details.append(("Transcription en annexe", f"traduite ({video.target_language})"))
    if template:
        details.append(("Modèle de compte-rendu", template.name))
    details.append(("Document généré le", datetime.now(timezone.utc).strftime("%d/%m/%Y à %H:%M UTC")))

    seconds = speaking_seconds(video)
    total = sum(seconds.values()) or 1.0
    labels = speaker_labels(video)
    speakers = [
        (speaker.label, _duration(seconds.get(speaker.id, 0.0)) if seconds.get(speaker.id, 0.0) >= 60
         else f"{int(seconds.get(speaker.id, 0.0))} s", f"{round(100 * seconds.get(speaker.id, 0.0) / total)} %")
        for speaker in video.speakers
    ]
    return Report(
        title=Path(video.original_filename).stem,
        details=details,
        speakers=speakers,
        summary=markdown_blocks(summary.content_markdown) if summary else [],
        chapters=[(timestamp(chapter.start_seconds), chapter.title) for chapter in video.chapters],
        transcript=(
            translated_rows(video.translated_text or "", list(labels.values()))
            if transcript == "translation"
            else [
                (timestamp(segment.start_seconds), labels.get(segment.speaker_id), segment.text)
                for segment in video.segments
            ] if transcript == "original" else []
        ),
    )


# --- DOCX ------------------------------------------------------------------------------

def render_docx(report: Report) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    document = Document()
    document.core_properties.title = report.title
    document.add_heading(report.title, level=0)

    table = document.add_table(rows=0, cols=2)
    table.style = "Light List"
    for label, value in report.details:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
        cells[0].paragraphs[0].runs[0].bold = True

    if report.speakers:
        document.add_heading("Intervenants", level=1)
        speakers = document.add_table(rows=1, cols=3)
        speakers.style = "Light List"
        for cell, title in zip(speakers.rows[0].cells, ("Intervenant", "Temps de parole", "Part")):
            cell.text = title
        for row in report.speakers:
            for cell, value in zip(speakers.add_row().cells, row):
                cell.text = value

    document.add_heading("Compte-rendu", level=1)
    if not report.summary:
        document.add_paragraph("Aucun résumé disponible.")
    for block in report.summary:
        if block.kind == "heading":
            document.add_heading(block.text, level=min(block.level + 1, 4))
            continue
        style = {"bullet": "List Bullet", "numbered": "List Number"}.get(block.kind)
        paragraph = document.add_paragraph(style=style)
        for text, bold in inline_runs(block.text):
            paragraph.add_run(text).bold = bold

    if report.chapters:
        document.add_heading("Chapitres", level=1)
        for time, title in report.chapters:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{time}  ").font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            paragraph.add_run(title)

    if report.transcript:
        document.add_heading("Transcription", level=1)
        for time, speaker, text in report.transcript:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            stamp = paragraph.add_run(f"[{time}] ")
            stamp.font.size = Pt(8)
            stamp.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            if speaker:
                paragraph.add_run(f"{speaker} : ").bold = True
            paragraph.add_run(text).font.size = Pt(10)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


# --- PDF -------------------------------------------------------------------------------

def _pdf_fonts() -> tuple[str, str]:
    """DejaVu covers every accent and symbol; Helvetica only Latin-1."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    regular, bold = FONT_DIR / "DejaVuSans.ttf", FONT_DIR / "DejaVuSans-Bold.ttf"
    if not (regular.is_file() and bold.is_file()):
        return "Helvetica", "Helvetica-Bold"
    if "DejaVuSans" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVuSans", str(regular)))
        pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", str(bold)))
        pdfmetrics.registerFontFamily("DejaVuSans", normal="DejaVuSans", bold="DejaVuSans-Bold")
    return "DejaVuSans", "DejaVuSans-Bold"


def _markup(text: str) -> str:
    """reportlab Paragraph mini-markup: escape, then **bold**."""
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return _BOLD.sub(r"<b>\1</b>", escaped)


def pdf_story(report: Report) -> list:
    """The report's flowables (separate from rendering, so tests can inspect them)."""
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, Spacer, Table, TableStyle

    regular, bold = _pdf_fonts()
    grey = colors.HexColor("#6B7280")
    body = ParagraphStyle("body", fontName=regular, fontSize=10, leading=14, spaceAfter=4)
    styles = {
        "title": ParagraphStyle("title", fontName=bold, fontSize=18, leading=22, spaceAfter=10),
        1: ParagraphStyle("h1", fontName=bold, fontSize=14, leading=18, spaceBefore=12, spaceAfter=6),
        2: ParagraphStyle("h2", fontName=bold, fontSize=12, leading=16, spaceBefore=8, spaceAfter=4),
        3: ParagraphStyle("h3", fontName=bold, fontSize=10.5, leading=14, spaceBefore=6, spaceAfter=3),
    }
    line = ParagraphStyle("line", parent=body, fontSize=9, leading=12, spaceAfter=1)

    def table(rows, header: bool = False):
        flowable = Table(rows, hAlign="LEFT", colWidths=None)
        commands = [
            ("FONTNAME", (0, 0), (-1, -1), regular), ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("TEXTCOLOR", (0, 0), (0, -1), grey), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")),
        ]
        if header:
            commands += [("FONTNAME", (0, 0), (-1, 0), bold), ("TEXTCOLOR", (0, 0), (-1, 0), colors.black)]
        flowable.setStyle(TableStyle(commands))
        return flowable

    story = [Paragraph(_markup(report.title), styles["title"]), table([list(row) for row in report.details]), Spacer(1, 6)]
    if report.speakers:
        story += [Paragraph("Intervenants", styles[1]),
                  table([["Intervenant", "Temps de parole", "Part"], *[list(row) for row in report.speakers]], header=True)]
    story.append(Paragraph("Compte-rendu", styles[1]))
    if not report.summary:
        story.append(Paragraph("Aucun résumé disponible.", body))
    items: list = []

    def flush_list() -> None:
        if items:
            # A copy: ListFlowable keeps the list it is given, and `items` is
            # cleared for the next list (every list of the PDF came out empty).
            story.append(ListFlowable(list(items), bulletType="bullet", bulletFontName=regular, leftIndent=12, bulletFontSize=8))
            items.clear()

    for block in report.summary:
        if block.kind in ("bullet", "numbered"):
            items.append(ListItem(Paragraph(_markup(block.text), body)))
            continue
        flush_list()
        if block.kind == "heading":
            story.append(Paragraph(_markup(block.text), styles[min(block.level + 1, 3)]))
        else:
            story.append(Paragraph(_markup(block.text), body))
    flush_list()

    if report.chapters:
        story.append(Paragraph("Chapitres", styles[1]))
        story += [Paragraph(f'<font color="#6B7280">{time}</font>  {_markup(title)}', body) for time, title in report.chapters]
    if report.transcript:
        story.append(Paragraph("Transcription", styles[1]))
        for time, speaker, text in report.transcript:
            who = f"<b>{_markup(speaker)} :</b> " if speaker else ""
            story.append(Paragraph(f'<font color="#6B7280" size="7.5">[{time}]</font> {who}{_markup(text)}', line))
    return story


def render_pdf(report: Report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate

    regular, _ = _pdf_fonts()
    grey = colors.HexColor("#6B7280")
    story = pdf_story(report)

    def footer(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont(regular, 8)
        canvas.setFillColor(grey)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"{report.title} · page {document.page}")
        canvas.restoreState()

    output = io.BytesIO()
    SimpleDocTemplate(
        output, pagesize=A4, title=report.title, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
